"""Rebuild every number in the JobSeeker-deaths draft post from source files.

Draft post: C:/dev/inthegame-blog/blog/2026-10-07-jobseeker-deaths/index.qmd
Findings:   analyses/jobseeker-deaths/FINDINGS.md

Inputs (gitignored, under data/):
  data/inputs/jobseeker-deaths/aihw-phe-354-income-support-mortality.xlsx   AIHW PHE 354 workbook
  data/inputs/jobseeker-deaths/dss-demographics/<YYYY-12>.xlsx, q2019-06.xlsx  DSS quarterly demographics
  data/extract/dss/FOI_19_20-073/.../p0029_t1.csv                          partial capacity 2007-2019 (pdfplumber)
  data/extract/vision/dss/16_17-018/t01_monthly_2012-2016.csv              Newstart monthly recipients (vision)
Outputs: data/outputs/jobseeker-deaths/*.csv

Every headline figure is asserted against the value quoted in the post, so a changed input fails
loudly instead of silently changing the story.
Run: python analyses/jobseeker-deaths/rebuild.py
"""
import csv, glob, os, re

import numpy as np
import openpyxl

ROOT = "C:/dev/foicommons/data"
INP = f"{ROOT}/inputs/jobseeker-deaths"
OUT = f"{ROOT}/outputs/jobseeker-deaths"
os.makedirs(OUT, exist_ok=True)


def write(name, rows):
    with open(f"{OUT}/{name}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(dict.fromkeys(k for r in rows for k in r)))
        w.writeheader()
        w.writerows(rows)
    print(f"  wrote {name}.csv ({len(rows)} rows)")


def near(a, b, tol, what):
    assert abs(a - b) <= tol, f"ANCHOR FAILED {what}: got {a}, expected {b}"


# ---------------------------------------------------------------- AIHW death rates
wb = openpyxl.load_workbook(f"{INP}/aihw-phe-354-income-support-mortality.xlsx", data_only=True)


def persons_rows(sheet):
    """Year rows under the 'Persons' block (or the whole sheet if there is none)."""
    rows = [[c for c in r if c is not None] for r in wb[sheet].iter_rows(values_only=True)]
    starts = [i for i, r in enumerate(rows) if r and r[0] == "Persons"]
    block = rows[starts[0] + 1:] if starts else rows
    out = {}
    for r in block:
        if r and isinstance(r[0], int) and 2011 < r[0] < 2023:
            out.setdefault(r[0], r)
    return out


GROUPS = {"all_recipients": "S1", "not_on_payment": "S2", "unemployment": "S3", "dsp": "S4", "carer": "S5", "parenting": "S6"}
aihw = []
for grp, t in GROUPS.items():
    crude, asr = persons_rows(f"Table {t}"), persons_rows(f"Table {t}.2")
    for y in sorted(crude):
        aihw.append({"group": grp, "year": y, "deaths": crude[y][-2], "crude_rate": crude[y][-1],
                     "asr": asr[y][-1], "population_est": round(crude[y][-2] / crude[y][-1] * 1e5)})
A = {(r["group"], r["year"]): r for r in aihw}
near(A[("unemployment", 2022)]["asr"], 325.7, 0.05, "JobSeeker ASR 2022")
near(A[("unemployment", 2020)]["asr"], 223.9, 0.05, "JobSeeker ASR 2020")
near(A[("not_on_payment", 2022)]["asr"], 66.8, 0.05, "non-recipient ASR 2022")
near(A[("unemployment", 2022)]["deaths"], 3985, 0, "JobSeeker deaths 2022")
near(A[("dsp", 2022)]["asr"], 1174, 0.5, "DSP ASR 2022")
print("AIHW rates ok")
write("aihw_rates", aihw)

# Jericho's counterfactuals
d22, r22, r20, nw22 = 3985, A[("unemployment", 2022)]["asr"], A[("unemployment", 2020)]["asr"], A[("not_on_payment", 2022)]["asr"]
near(round(d22 * (1 - r20 / r22)), 1246, 1, "Jericho 1,246")
near(round(d22 * nw22 / r22), 817, 1, "Jericho 817")

# Double counting: sum of payment-group deaths minus all-recipient deaths
dbl = []
for y in range(2012, 2023):
    s = sum(A[(g, y)]["deaths"] for g in ("unemployment", "dsp", "carer", "parenting"))
    over = s - A[("all_recipients", y)]["deaths"]
    dbl.append({"year": y, "jobseeker_deaths": A[("unemployment", y)]["deaths"], "in_2plus_groups": over,
                "max_share": round(over / A[("unemployment", y)]["deaths"], 3)})
near([r for r in dbl if r["year"] == 2022][0]["in_2plus_groups"], 706, 0, "double count 2022")
write("double_counting", dbl)

# ---------------------------------------------------------------- DSS partial capacity (December)
def dss_partial(path):
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = [s for s in w.sheetnames if s.strip() == "Partial Capacity to Work"]
    if not sh:
        return None
    for r in w[sh[0]].iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        if v and isinstance(v[0], str) and re.match(r"(Newstart|JobSeeker)", v[0].strip()):
            return int(v[1]), int(v[2]), int(v[3])
    return None


# Revised-definition back-series (June 2019 release) for Dec 2014-2018
back = {}
w = openpyxl.load_workbook(f"{INP}/dss-demographics/q2019-06.xlsx", read_only=True, data_only=True)
sh = [s for s in w.sheetnames if "apacity" in s][0]
for r in w[sh].iter_rows(values_only=True):
    v = [c for c in r if c is not None]
    if len(v) >= 5 and str(v[1]).startswith("Newstart") and hasattr(v[0], "month") and v[0].month == 12:
        back[v[0].year] = (int(v[2]), int(v[3]), int(v[4]))

pc = []
for f in sorted(glob.glob(f"{INP}/dss-demographics/20*-12.xlsx")):
    y = int(os.path.basename(f)[:4])
    vals = back.get(y) if y < 2019 else dss_partial(f)
    if vals:
        full, partial, total = vals
        pc.append({"year": y, "full_capacity": full, "partial": partial, "total": total,
                   "partial_share": round(partial / total, 4), "basis": "revised back-series" if y < 2019 else "published"})
P = {r["year"]: r for r in pc}
near(P[2014]["partial"], 195054, 0, "partial Dec 2014 (revised)")
near(P[2019]["partial_share"], 0.4281, 0.0001, "partial share Dec 2019")
near(P[2020]["partial_share"], 0.2905, 0.0001, "partial share Dec 2020")
print("DSS partial capacity ok")
write("partial_capacity_dec", pc)

# ---------------------------------------------------------------- straight-line estimates
yrs = list(range(2014, 2023))
s = np.array([P[y]["partial_share"] for y in yrs])


def fit(x, y, label):
    X = np.c_[np.ones_like(x), x]
    b = np.linalg.lstsq(X, y, rcond=None)[0]
    n = len(x)
    sig2 = ((y - X @ b) ** 2).sum() / (n - 2)
    cov = sig2 * np.linalg.inv(X.T @ X)
    tcrit = {7: 2.365, 6: 2.447}[n - 2]
    full, partial = b[0], b[0] + b[1]
    se_f, se_p = np.sqrt(cov[0, 0]), np.sqrt(cov[0, 0] + cov[1, 1] + 2 * cov[0, 1])
    return {"version": label, "n_years": n, "full": round(full, 1), "full_lo": round(full - tcrit * se_f, 1),
            "full_hi": round(full + tcrit * se_f, 1), "partial": round(partial, 1), "partial_pm": round(tcrit * se_p, 1),
            "corr": round(np.corrcoef(x, y)[0, 1], 3)}


asr = np.array([A[("unemployment", y)]["asr"] for y in yrs])
crude = np.array([A[("unemployment", y)]["crude_rate"] for y in yrs])
keep = np.array([y != 2020 for y in yrs])
est = [fit(s, asr, "age-standardised, all 9 years"), fit(s[keep], asr[keep], "age-standardised, without 2020"),
       fit(s, crude, "crude, all 9 years")]
near(est[0]["full"], 52, 1, "full-capacity estimate"); near(est[0]["partial"], 696, 1, "partial estimate")
near(est[1]["full"], 68, 1, "full without 2020"); near(est[0]["corr"], 0.976, 0.002, "correlation")
print("estimates ok")
write("estimates", est)

# Out-of-sample: fit 2014-2019, predict 2020-2022 (all causes and deaths of despair)
b1, a1 = np.polyfit(s[:6], asr[:6], 1)
oos = [{"year": y, "asr_actual": A[("unemployment", y)]["asr"], "asr_predicted": round(a1 + b1 * P[y]["partial_share"], 1)} for y in (2020, 2021, 2022)]
near(oos[0]["asr_predicted"], 244, 1, "2020 prediction")
write("out_of_sample", oos)


def despair(sheet):
    rows = [[c for c in r if c is not None] for r in wb[sheet].iter_rows(values_only=True)]
    i = [k for k, r in enumerate(rows) if r and str(r[0]).strip() == "Total"][0]
    return {int(r[0]): float(r[-1]) for r in rows[i + 1:i + 12] if r and str(r[0]).isdigit()}


dj, dn = despair("Table S13"), despair("Table S12")
dsp_rows = []
bd, ad = np.polyfit(s[:6], np.array([dj[y] for y in range(2014, 2020)]), 1)
bc, ac = np.polyfit(s[:6], crude[:6], 1)
for y in range(2014, 2023):
    row = {"year": y, "jobseeker_despair_rate": dj[y], "non_recipient_despair_rate": dn[y]}
    if y >= 2020:
        pd_, pc_ = ad + bd * P[y]["partial_share"], ac + bc * P[y]["partial_share"]
        row.update(predicted=round(pd_, 1), gap=round(dj[y] / pd_ - 1, 3),
                   all_cause_gap=round(A[("unemployment", y)]["crude_rate"] / pc_ - 1, 3))
    dsp_rows.append(row)
near(dj[2019], 116.4, 0.05, "despair 2019"); near(dsp_rows[6]["gap"], -0.125, 0.002, "despair gap 2020")
write("deaths_of_despair", dsp_rows)

# ---------------------------------------------------------------- FOI-derived series
foi73 = glob.glob(f"{ROOT}/extract/dss/FOI_19_20-073/*/p0029_t1.csv")[0]
with open(foi73, encoding="utf-8") as fh:
    r73 = list(csv.reader(fh))
pc0719 = [{"as_at": r[0], "year": int(r[0][-4:]), "partial_capacity_total": int(r[-1])} for r in r73[1:] if r and r[0][:2].isdigit()]
near([r for r in pc0719 if r["year"] == 2019][0]["partial_capacity_total"], 289489, 0, "partial June 2019 (FOI 19/20-073)")
write("partial_capacity_june_2007_2019", pc0719)

with open(f"{ROOT}/extract/vision/dss/16_17-018/t01_monthly_2012-2016.csv", encoding="utf-8") as fh:
    nsa = {r["month"]: int(r["newstart_recipients"]) for r in csv.DictReader(fh)}
J = {r["year"]: r["partial_capacity_total"] for r in pc0719}
share_jun = [{"june": y, "partial": J[y], "newstart_recipients": nsa[m], "share": round(J[y] / nsa[m], 3)}
             for y, m in ((2013, "Jun-13"), (2014, "Jun-14"), (2015, "Jun-15"), (2016, "Jun-16"))]
near(share_jun[0]["share"], 0.229, 0.001, "June 2013 share")
write("partial_share_june_2013_2016", share_jun)

# DSS FOI 2122-068: cancellations due to death (transcribed from the released PDF, all ages)
on_payment = {2019: 1903, 2020: 2567, 2021: 2223}
write("on_payment_at_death", [{"year": y, "dss_cancelled_due_to_death": v, "aihw_jobseeker_deaths": A[("unemployment", y)]["deaths"],
                              "share": round(v / A[("unemployment", y)]["deaths"], 3)} for y, v in on_payment.items()])

print("ALL ANCHORS PASS")
